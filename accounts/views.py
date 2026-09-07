from django.contrib import messages
from django.contrib.auth import get_user_model, login
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.views.decorators.http import require_POST

from .forms import ProfileForm, SignupForm
from .models import FriendRequest, friends_of


def signup(request):
    if request.user.is_authenticated:
        return redirect("dashboard:home")
    if request.method == "POST":
        form = SignupForm(request.POST)
        if form.is_valid():
            user = form.save()
            login(request, user)
            return redirect("dashboard:home")
    else:
        form = SignupForm()
    return render(request, "accounts/signup.html", {"form": form})


@login_required
def profile(request):
    if request.method == "POST":
        form = ProfileForm(request.POST, request.FILES, instance=request.user)
        if form.is_valid():
            form.save()
            messages.success(request, "Perfil actualizado.")
            return redirect("accounts:profile")
    else:
        form = ProfileForm(instance=request.user)
    return render(request, "accounts/profile.html", {"form": form})


@login_required
def export_excel(request):
    """Copia de seguridad personal en Excel: todo lo que ha registrado el
    usuario (entrenamientos con sus series, récords personales, peso
    corporal, objetivos) -- para no depender solo de la base de datos si
    algún día quiere llevarse sus datos a otro sitio o guardarlos aparte."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    from training.models import BodyWeightEntry, Goal, PersonalRecord, SetEntry, Workout

    user = request.user
    wb = Workbook()

    ws_workouts = wb.active
    ws_workouts.title = "Entrenamientos"
    ws_workouts.append(["Fecha", "Nombre", "Estado", "Duración (min)", "Energía", "RPE sesión", "Notas"])
    for w in Workout.objects.filter(user=user).order_by("-date"):
        ws_workouts.append([
            w.date.strftime("%d/%m/%Y"), w.name, w.get_status_display(),
            w.duration_minutes, w.get_energy_level_display() if w.energy_level else "",
            float(w.rpe_session) if w.rpe_session is not None else "", w.notes,
        ])

    ws_sets = wb.create_sheet("Series")
    ws_sets.append(["Fecha", "Entrenamiento", "Ejercicio", "Serie", "Tipo", "Peso (kg)", "Reps", "RPE", "1RM estimado (kg)"])
    sets = (
        SetEntry.objects.filter(workout_exercise__workout__user=user)
        .select_related("workout_exercise__workout", "workout_exercise__exercise")
        .order_by("-workout_exercise__workout__date", "workout_exercise__workout_id", "workout_exercise__order", "set_number")
    )
    for s in sets:
        workout = s.workout_exercise.workout
        ws_sets.append([
            workout.date.strftime("%d/%m/%Y"), workout.name, s.workout_exercise.exercise.name, s.set_number,
            s.get_set_type_display(), float(s.weight_kg) if s.weight_kg is not None else "",
            s.reps_performed or "", float(s.rpe) if s.rpe is not None else "",
            float(s.estimated_1rm_kg) if s.estimated_1rm_kg is not None else "",
        ])

    ws_prs = wb.create_sheet("Récords personales")
    ws_prs.append(["Ejercicio", "Tipo", "Valor (kg)", "Reps", "Conseguido"])
    for pr in PersonalRecord.objects.filter(user=user).select_related("exercise").order_by("exercise__name", "kind"):
        ws_prs.append([
            pr.exercise.name, pr.get_kind_display(), float(pr.value_kg) if pr.value_kg is not None else "",
            pr.reps or "", pr.achieved_at.strftime("%d/%m/%Y %H:%M"),
        ])

    ws_bw = wb.create_sheet("Peso corporal")
    ws_bw.append(["Fecha", "Peso", "Unidad", "% grasa corporal", "Notas"])
    for entry in BodyWeightEntry.objects.filter(user=user).order_by("-date"):
        ws_bw.append([
            entry.date.strftime("%d/%m/%Y"), float(entry.weight_display), entry.input_unit,
            float(entry.body_fat_pct) if entry.body_fat_pct is not None else "", entry.notes,
        ])

    ws_goals = wb.create_sheet("Objetivos")
    ws_goals.append(["Título", "Tipo", "Fecha objetivo", "Conseguido", "Notas"])
    for goal in Goal.objects.filter(user=user).order_by("achieved", "target_date"):
        ws_goals.append([
            goal.title, goal.get_kind_display(),
            goal.target_date.strftime("%d/%m/%Y") if goal.target_date else "",
            "Sí" if goal.achieved else "No", goal.notes,
        ])

    for ws in wb.worksheets:
        for col_idx, column_cells in enumerate(ws.columns, start=1):
            max_len = max((len(str(cell.value)) for cell in column_cells if cell.value is not None), default=10)
            ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 60)

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    filename = f"tonnage_backup_{user.username}_{timezone.now():%Y%m%d_%H%M}.xlsx"
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response


@login_required
def friends_list(request):
    User = get_user_model()
    query = request.GET.get("q", "").strip()
    results = []
    if query:
        results = (
            User.objects.filter(username__icontains=query)
            .exclude(pk=request.user.pk)
            .exclude(pk__in=[f.pk for f in friends_of(request.user)])[:10]
        )

    friends = friends_of(request.user)
    incoming = FriendRequest.objects.filter(to_user=request.user, accepted=False).select_related("from_user")
    outgoing = FriendRequest.objects.filter(from_user=request.user, accepted=False).select_related("to_user")

    return render(request, "accounts/friends_list.html", {
        "query": query, "results": results,
        "friends": friends, "incoming": incoming, "outgoing": outgoing,
    })


@login_required
@require_POST
def friend_request_send(request, username):
    User = get_user_model()
    other = get_object_or_404(User, username=username)
    if other.pk != request.user.pk:
        existing = FriendRequest.objects.filter(from_user=other, to_user=request.user, accepted=False).first()
        if existing:
            # El otro ya te había pedido amistad a ti -- aceptarla directamente
            # en vez de dejar dos solicitudes cruzadas sin resolver.
            existing.accepted = True
            existing.save(update_fields=["accepted"])
            messages.success(request, f"Ahora eres amigo de {other.username}.")
        else:
            FriendRequest.objects.get_or_create(from_user=request.user, to_user=other)
            messages.success(request, f"Solicitud enviada a {other.username}.")
    return redirect("accounts:friends")


@login_required
@require_POST
def friend_request_accept(request, pk):
    req = get_object_or_404(FriendRequest, pk=pk, to_user=request.user, accepted=False)
    req.accepted = True
    req.save(update_fields=["accepted"])
    messages.success(request, f"Ahora eres amigo de {req.from_user.username}.")
    return redirect("accounts:friends")


@login_required
@require_POST
def friend_request_decline(request, pk):
    get_object_or_404(FriendRequest, pk=pk, to_user=request.user, accepted=False).delete()
    return redirect("accounts:friends")


@login_required
@require_POST
def friend_request_withdraw(request, pk):
    get_object_or_404(FriendRequest, pk=pk, from_user=request.user, accepted=False).delete()
    return redirect("accounts:friends")
