from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse

User = get_user_model()


class ProfileTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="lifter", password="pass12345")
        self.client.force_login(self.user)

    def test_cambiar_unidad_preferida(self):
        response = self.client.get(reverse("accounts:profile"))
        self.assertEqual(response.status_code, 200)

        self.client.post(reverse("accounts:profile"), {
            "unit_preference": "lb", "height_cm": "", "birth_date": "",
        })
        self.user.refresh_from_db()
        self.assertEqual(self.user.unit_preference, "lb")

    def test_perfil_requiere_login(self):
        self.client.logout()
        response = self.client.get(reverse("accounts:profile"))
        self.assertEqual(response.status_code, 302)


class SignupTests(TestCase):
    def test_registro_crea_usuario_y_entra(self):
        response = self.client.post(reverse("accounts:signup"), {
            "username": "nuevo", "password1": "unaContrasenaSegura123", "password2": "unaContrasenaSegura123",
        })
        self.assertEqual(User.objects.filter(username="nuevo").count(), 1)
        self.assertRedirects(response, reverse("dashboard:home"))


class ExportExcelTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_user(username="lifter", password="pass12345")
        self.client.force_login(self.user)

    def test_exportar_devuelve_un_xlsx(self):
        response = self.client.get(reverse("accounts:export-excel"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("attachment", response["Content-Disposition"])

    def test_exportar_incluye_los_entrenamientos_del_usuario(self):
        from decimal import Decimal

        from training.models import Exercise, MuscleGroup, SetEntry, Workout, WorkoutExercise

        chest = MuscleGroup.objects.create(name="Pecho", region="upper")
        bench = Exercise.objects.create(name="Press banca", primary_muscle=chest, equipment="barbell")
        workout = Workout.objects.create(user=self.user, name="Empuje", start_time=Workout.start_time.field.default())
        we = WorkoutExercise.objects.create(workout=workout, exercise=bench)
        SetEntry.objects.create(workout_exercise=we, set_number=1, weight_kg=Decimal("80"), reps_performed=8)

        response = self.client.get(reverse("accounts:export-excel"))

        import io

        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(response.content))
        self.assertIn("Entrenamientos", wb.sheetnames)
        self.assertIn("Series", wb.sheetnames)
        rows = list(wb["Entrenamientos"].iter_rows(values_only=True))
        self.assertEqual(rows[1][1], "Empuje")

    def test_exportar_requiere_login(self):
        self.client.logout()
        response = self.client.get(reverse("accounts:export-excel"))
        self.assertEqual(response.status_code, 302)


class FriendshipHelpersTests(TestCase):
    def setUp(self):
        self.ana = User.objects.create_user(username="ana", password="pass12345")
        self.bea = User.objects.create_user(username="bea", password="pass12345")

    def test_sin_solicitud_no_son_amigos(self):
        from .models import are_friends, friendship_status

        self.assertFalse(are_friends(self.ana, self.bea))
        self.assertEqual(friendship_status(self.ana, self.bea), "none")

    def test_solicitud_aceptada_los_hace_amigos(self):
        from .models import FriendRequest, are_friends, friendship_status

        FriendRequest.objects.create(from_user=self.ana, to_user=self.bea, accepted=True)
        self.assertTrue(are_friends(self.ana, self.bea))
        self.assertTrue(are_friends(self.bea, self.ana))
        self.assertEqual(friendship_status(self.ana, self.bea), "friends")

    def test_solicitud_pendiente_se_ve_desde_ambos_lados(self):
        from .models import FriendRequest, friendship_status

        FriendRequest.objects.create(from_user=self.ana, to_user=self.bea)
        self.assertEqual(friendship_status(self.ana, self.bea), "pending_outgoing")
        self.assertEqual(friendship_status(self.bea, self.ana), "pending_incoming")


class FriendRequestViewsTests(TestCase):
    def setUp(self):
        self.ana = User.objects.create_user(username="ana", password="pass12345")
        self.bea = User.objects.create_user(username="bea", password="pass12345")
        self.client.force_login(self.ana)

    def test_enviar_solicitud(self):
        from .models import FriendRequest

        self.client.post(reverse("accounts:friend-request-send", args=["bea"]))
        self.assertTrue(FriendRequest.objects.filter(from_user=self.ana, to_user=self.bea, accepted=False).exists())

    def test_solicitud_cruzada_se_acepta_directamente(self):
        from .models import FriendRequest, are_friends

        FriendRequest.objects.create(from_user=self.bea, to_user=self.ana, accepted=False)
        self.client.post(reverse("accounts:friend-request-send", args=["bea"]))
        self.assertTrue(are_friends(self.ana, self.bea))

    def test_aceptar_solicitud(self):
        from .models import FriendRequest, are_friends

        req = FriendRequest.objects.create(from_user=self.bea, to_user=self.ana)
        self.client.post(reverse("accounts:friend-request-accept", args=[req.pk]))
        self.assertTrue(are_friends(self.ana, self.bea))

    def test_no_se_puede_aceptar_la_solicitud_de_otro(self):
        from .models import FriendRequest

        carla = User.objects.create_user(username="carla", password="pass12345")
        req = FriendRequest.objects.create(from_user=self.bea, to_user=carla)
        response = self.client.post(reverse("accounts:friend-request-accept", args=[req.pk]))
        self.assertEqual(response.status_code, 404)

    def test_rechazar_solicitud_la_borra(self):
        from .models import FriendRequest

        req = FriendRequest.objects.create(from_user=self.bea, to_user=self.ana)
        self.client.post(reverse("accounts:friend-request-decline", args=[req.pk]))
        self.assertFalse(FriendRequest.objects.filter(pk=req.pk).exists())

    def test_pagina_de_amigos_requiere_login(self):
        self.client.logout()
        response = self.client.get(reverse("accounts:friends"))
        self.assertEqual(response.status_code, 302)
